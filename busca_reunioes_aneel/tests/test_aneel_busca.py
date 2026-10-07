import io
from datetime import date
from pathlib import Path

import pytest
import yaml

from aneel_busca.busca import Buscador
from aneel_busca.cliente import Resposta
from aneel_busca.extracao import extrair_texto
from aneel_busca.processamento import Varredura
from aneel_busca.relatorio import gravar_excel
from aneel_busca.site_aneel import eh_desafio_cloudflare, ler_detalhe, ler_listagem, url_listagem

FIXTURES = Path(__file__).parent / "fixtures"
CONFIG = Path(__file__).parent.parent / "config.yaml"


def config_teste(**extra):
    config = yaml.safe_load(CONFIG.read_text(encoding="utf-8"))
    config["processos"] = ["48500.001234/2025-11", "002222/2026"]
    config["palavras_chave"] = ["revisão tarifária"]
    config.update(extra)
    return config


def pdf_minimo(texto: str) -> bytes:
    """Gera um PDF simples (uma página, fonte Helvetica) sem dependências externas."""
    linhas = texto.split("\n")
    conteudo = "BT /F1 11 Tf 50 800 Td 14 TL " + " ".join(
        "(" + l.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") '" for l in linhas) + " ET"
    objetos = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(conteudo.encode('latin-1'))} >>\nstream\n{conteudo}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    saida = io.BytesIO()
    saida.write(b"%PDF-1.4\n")
    posicoes = []
    for i, obj in enumerate(objetos, start=1):
        posicoes.append(saida.tell())
        saida.write(f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1"))
    xref = saida.tell()
    saida.write(f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode())
    for p in posicoes:
        saida.write(f"{p:010d} 00000 n \n".encode())
    saida.write(f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return saida.getvalue()


# ---------------------------------------------------------------- listagens reais salvas do site
def test_listagem_pautas_e_atas():
    listagem = ler_listagem((FIXTURES / "lista_425_pagina1.html").read_bytes())
    assert len(listagem.reunioes) == 15
    primeira = listagem.reunioes[0]
    assert primeira.data == date(2026, 10, 6)
    assert primeira.titulo == "PAUTA/ATA DA 20ª REUNIÃO PÚBLICA ORDINÁRIA DA DIRETORIA DE 2026 (PRÉVIA)"
    assert primeira.previa
    assert primeira.id_noticia == "14817"
    assert "dsp_detalheNoticia.cfm?idNoticia=14817&idAreaNoticia=425" in primeira.url
    assert not listagem.reunioes[1].previa
    assert listagem.proxima_pagina.endswith("page=2&txtPesq=&anoPesq=&mesPesq=")


def test_listagem_distribuicao_pagina2():
    listagem = ler_listagem((FIXTURES / "lista_424_pagina2.html").read_bytes())
    assert len(listagem.reunioes) == 15
    assert listagem.reunioes[0].data == date(2026, 7, 13)
    assert "DISTRIBUIÇÃO DE PROCESSOS" in listagem.reunioes[0].titulo
    assert "page=3" in listagem.proxima_pagina  # não confunde com o link "15 Anteriores"


def test_url_listagem():
    assert url_listagem(425, 2).endswith("idAreaNoticia=425&idarea=425&page=2&txtPesq=&anoPesq=&mesPesq=")


def test_detecta_cloudflare():
    assert eh_desafio_cloudflare(403, b"<title>Just a moment...</title><script>window._cf_chl_opt={}</script>")
    pagina_normal = (FIXTURES / "lista_425_pagina1.html").read_bytes()  # tem o beacon do Cloudflare
    assert not eh_desafio_cloudflare(200, pagina_normal)


# ---------------------------------------------------------------- busca
TEXTO_PAUTA = """PAUTA DA 20ª REUNIÃO PÚBLICA ORDINÁRIA
1. Processo: 48500.001234/2025-11
Interessada: ENERGISA SERGIPE - DISTRIBUIDORA DE ENERGIA S.A.
Assunto: Revisao Tarifaria Periodica.
2. Processo: 48500.009999/2026-01
Interessada: Companhia Paulista de Força e Luz
Assunto: Recurso administrativo da EMT contra auto de infração. Sistemas de emissão (ems) não contam.
3. Processo 48500002222202677 - Grupo Energisa: fiscalização.
"""


def test_busca_processos_empresas_e_palavras():
    buscador = Buscador.de_config(config_teste())
    ocorrencias = buscador.buscar(TEXTO_PAUTA)
    por_termo = {}
    for o in ocorrencias:
        por_termo.setdefault(o.termo, []).append(o)

    assert [o.encontrado for o in por_termo["48500.001234/2025-11"]] == ["48500.001234/2025-11"]
    assert [o.encontrado for o in por_termo["0022222026"]] == ["48500002222202677"]  # trecho parcial
    assert [o.encontrado for o in por_termo["Energisa Sergipe (ESE)"]] == ["ENERGISA SERGIPE"]
    assert [o.encontrado for o in por_termo["Energisa Mato Grosso (EMT)"]] == ["EMT"]
    assert "Energisa Mato Grosso do Sul (EMS)" not in por_termo  # "ems" minúsculo não é a sigla
    # "Energisa" genérico só onde não há empresa específica
    assert [o.encontrado for o in por_termo["Grupo Energisa (genérico)"]] == ["Energisa"]
    # sem acento no texto, com acento no termo
    assert [o.encontrado for o in por_termo["revisão tarifária"]] == ["Revisao Tarifaria"]

    sergipe = por_termo["Energisa Sergipe (ESE)"][0]
    assert sergipe.item == "1. Processo: 48500.001234/2025-11"
    assert "[[ENERGISA SERGIPE]]" in sergipe.trecho
    assert por_termo["Energisa Mato Grosso (EMT)"][0].item.startswith("2. Processo: 48500.009999/2026-01")


def test_mato_grosso_do_sul_nao_conta_como_mato_grosso():
    buscador = Buscador.de_config(config_teste())
    termos = [o.termo for o in buscador.buscar("Recurso da Energisa Mato Grosso do Sul.")]
    assert termos == ["Energisa Mato Grosso do Sul (EMS)"]


def test_processos_citados():
    buscador = Buscador.de_config(config_teste())
    assert buscador.processos_citados(TEXTO_PAUTA) == [
        "48500.001234/2025-11", "48500.009999/2026-01", "48500.002222/2026-77"]


# ---------------------------------------------------------------- extração de arquivos
def test_extrai_pdf():
    tipo, texto = extrair_texto(pdf_minimo("Processo 48500.001234/2025-11\nEnergisa Tocantins"))
    assert tipo == "pdf"
    assert "48500.001234/2025-11" in texto and "Energisa Tocantins" in texto


def test_extrai_docx():
    import docx

    documento = docx.Document()
    documento.add_paragraph("Interessada: Energisa Rondônia")
    tabela = documento.add_table(rows=1, cols=2)
    tabela.cell(0, 0).text = "48500.001234/2025-11"
    tabela.cell(0, 1).text = "EAC"
    saida = io.BytesIO()
    documento.save(saida)
    tipo, texto = extrair_texto(saida.getvalue(), nome="pauta.docx")
    assert tipo == "docx"
    assert "Energisa Rondônia" in texto and "48500.001234/2025-11 | EAC" in texto


def test_html_windows_1252():
    html = "<html><head><meta charset='windows-1252'></head><body>Reunião Pública – Paraíba</body></html>"
    tipo, texto = extrair_texto(html.encode("cp1252"))
    assert tipo == "html" and texto == "Reunião Pública – Paraíba"


# ---------------------------------------------------------------- fluxo completo com site simulado
DETALHE = """<html><head><meta charset="windows-1252"><title>Pauta</title></head><body>
<h1>PAUTA/ATA DA 19ª REUNIÃO PÚBLICA ORDINÁRIA DA DIRETORIA DE 2026</h1>
<p>Item 1 - Processo 48500.001234/2025-11 - Energisa Paraíba.</p>
<a href="/aplicacoes_liferay/noticias_area/arquivo.cfm?id=1">Pauta completa</a>
<a href="https://www2.aneel.gov.br/documents/123/ata.pdf">Ata (PDF)</a>
<a href="mailto:reuniaodir@aneel.gov.br">Fale conosco</a>
<a href="dsp_listarNoticias.cfm?idAreaNoticia=425">Voltar</a>
</body></html>""".encode("cp1252")


class ClienteFalso:
    def __init__(self, paginas):
        self.paginas = paginas
        self.pedidos = []

    def obter(self, url, usar_cache=True):
        self.pedidos.append((url, usar_cache))
        for trecho, (conteudo, tipo) in self.paginas.items():
            if trecho in url:
                return Resposta(url, 200, conteudo, tipo)
        return Resposta(url, 404, b"", "text/html")


def test_detalhe_identifica_anexos():
    detalhe = ler_detalhe(DETALHE, "https://www2.aneel.gov.br/aplicacoes_liferay/noticias_area/dsp_detalheNoticia.cfm")
    assert [nome for nome, _ in detalhe.anexos] == ["Pauta completa", "Ata (PDF)"]
    assert detalhe.anexos[0][1] == "https://www2.aneel.gov.br/aplicacoes_liferay/noticias_area/arquivo.cfm?id=1"
    assert "Energisa Paraíba" in detalhe.texto


def test_varredura_online_simulada(tmp_path):
    lista = (FIXTURES / "lista_425_pagina1.html").read_bytes()
    cliente = ClienteFalso({
        "dsp_listarNoticias": (lista, "text/html"),
        "idNoticia=14810": (DETALHE, "text/html"),  # 19ª reunião, 22/09/2026
        "arquivo.cfm?id=1": (pdf_minimo("Interessada: Energisa Sul-Sudeste"), "application/pdf"),
        "ata.pdf": (pdf_minimo("Ata: aprovado o pleito da EPB"), "application/octet-stream"),
    })
    varredura = Varredura(Buscador.de_config(config_teste()))
    resultado = varredura.varrer_online(cliente, {425: "Pautas e Atas"}, desde=date(2026, 9, 20),
                                        ate=date(2026, 9, 30), max_paginas=3)

    reunioes = {d.reuniao for d in resultado.documentos}
    # "PAUTAATA" está escrito assim mesmo no site
    assert reunioes == {"PAUTAATA DO 17º CIRCUITO DELIBERATIVO PÚBLICO ORDINÁRIO DA DIRETORIA DE 2026.",
                        "PAUTA\\ATA DA 19ª REUNIÃO PÚBLICA ORDINÁRIA DA DIRETORIA DE 2026"}
    # a listagem tem reuniões anteriores a "desde": não deve buscar a página 2
    assert sum("dsp_listarNoticias" in url for url, _ in cliente.pedidos) == 1
    achados = {(o.documento, o.termo) for o in resultado.ocorrencias}
    assert ("Página da reunião", "48500.001234/2025-11") in achados
    assert ("Página da reunião", "Energisa Paraíba (EPB)") in achados
    assert ("Pauta completa", "Energisa Sul-Sudeste (ESS)") in achados
    assert ("Ata (PDF)", "Energisa Paraíba (EPB)") in achados
    # a reunião do circuito (sem página simulada) aparece como erro, sem derrubar a varredura
    assert any(d.situacao.startswith("ERRO") for d in resultado.documentos)

    xlsx = tmp_path / "saida.xlsx"
    gravar_excel(xlsx, resultado.ocorrencias, resultado.documentos)
    import openpyxl
    livro = openpyxl.load_workbook(xlsx)
    assert livro.sheetnames == ["Resumo", "Ocorrências", "Documentos analisados"]
    assert livro["Ocorrências"].max_row == len(resultado.ocorrencias) + 1


def test_previa_nao_usa_cache():
    lista = (FIXTURES / "lista_425_pagina1.html").read_bytes()
    cliente = ClienteFalso({"dsp_listarNoticias": (lista, "text/html"), "idNoticia=14817": (DETALHE, "text/html")})
    Varredura(Buscador.de_config(config_teste())).varrer_online(
        cliente, {425: "Pautas e Atas"}, desde=date(2026, 10, 1), ate=None, max_paginas=1)
    pedidos = dict(cliente.pedidos)
    assert pedidos[next(u for u in pedidos if "idNoticia=14817" in u)] is False


def test_varredura_local(tmp_path):
    (tmp_path / "pauta.pdf").write_bytes(pdf_minimo("22/09/2026\nProcesso 48500.001234/2025-11 - ERO"))
    (tmp_path / "quebrado.pdf").write_bytes(b"%PDF-1.4 lixo")
    (tmp_path / "pagina_files").mkdir()
    (tmp_path / "pagina_files" / "x.html").write_text("<html>Energisa</html>")
    resultado = Varredura(Buscador.de_config(config_teste())).varrer_pasta(tmp_path)
    assert {o.termo for o in resultado.ocorrencias} == {"48500.001234/2025-11", "Energisa Rondônia (ERO)"}
    assert resultado.ocorrencias[0].data == "22/09/2026"
    assert [d.situacao.startswith("ERRO") for d in resultado.documentos] == [True, False] or \
           [d.situacao.startswith("ERRO") for d in resultado.documentos] == [False, True]


def test_cli_local(tmp_path, capsys):
    from aneel_busca.__main__ import main

    pasta = tmp_path / "docs"
    pasta.mkdir()
    (pasta / "ata.pdf").write_bytes(pdf_minimo("Energisa Acre"))
    codigo = main(["--config", str(CONFIG), "--saida", str(tmp_path / "res"), "local", str(pasta)])
    assert codigo == 0
    assert "Energisa Acre (EAC): 1 reunião" in capsys.readouterr().out
    assert len(list((tmp_path / "res").glob("busca_aneel_*.xlsx"))) == 1


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
