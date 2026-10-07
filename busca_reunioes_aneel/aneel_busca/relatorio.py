"""Geração dos relatórios (Excel e CSV)."""

from __future__ import annotations

import csv
from dataclasses import asdict, dataclass
from pathlib import Path


@dataclass
class LinhaOcorrencia:
    data: str
    area: str
    reuniao: str
    url_reuniao: str
    documento: str
    url_documento: str
    categoria: str
    termo: str
    encontrado: str
    item: str
    trecho: str


@dataclass
class LinhaDocumento:
    data: str
    area: str
    reuniao: str
    documento: str
    url_documento: str
    tipo: str
    caracteres: int
    ocorrencias: int
    processos_citados: str
    situacao: str


COLUNAS_OCORRENCIAS = {
    "data": "Data", "area": "Área", "reuniao": "Reunião", "url_reuniao": "Link da reunião",
    "documento": "Documento", "url_documento": "Link do documento", "categoria": "Categoria",
    "termo": "Termo buscado", "encontrado": "Texto encontrado", "item": "Item da pauta (aprox.)",
    "trecho": "Trecho",
}
COLUNAS_DOCUMENTOS = {
    "data": "Data", "area": "Área", "reuniao": "Reunião", "documento": "Documento",
    "url_documento": "Link do documento", "tipo": "Tipo", "caracteres": "Caracteres lidos",
    "ocorrencias": "Ocorrências", "processos_citados": "Processos citados", "situacao": "Situação",
}


def gravar_csv(caminho: Path, linhas: list, colunas: dict[str, str]) -> None:
    with caminho.open("w", newline="", encoding="utf-8-sig") as arquivo:  # utf-8-sig: abre certo no Excel
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(colunas.values())
        for linha in linhas:
            d = asdict(linha)
            escritor.writerow([d[c] for c in colunas])


def gravar_excel(caminho: Path, ocorrencias: list[LinhaOcorrencia], documentos: list[LinhaDocumento]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    resumo = livro.active
    resumo.title = "Resumo"
    _aba_resumo(resumo, ocorrencias, documentos)

    larguras_oc = {"data": 11, "area": 16, "reuniao": 45, "url_reuniao": 30, "documento": 30, "url_documento": 30,
                   "categoria": 14, "termo": 32, "encontrado": 24, "item": 40, "trecho": 90}
    larguras_doc = {"data": 11, "area": 16, "reuniao": 45, "documento": 35, "url_documento": 35, "tipo": 8,
                    "caracteres": 12, "ocorrencias": 12, "processos_citados": 60, "situacao": 40}
    for titulo, linhas, colunas, larguras in (
        ("Ocorrências", ocorrencias, COLUNAS_OCORRENCIAS, larguras_oc),
        ("Documentos analisados", documentos, COLUNAS_DOCUMENTOS, larguras_doc),
    ):
        aba = livro.create_sheet(titulo)
        aba.append(list(colunas.values()))
        for linha in linhas:
            d = asdict(linha)
            aba.append([_celula(d[c]) for c in colunas])
        for i, chave in enumerate(colunas, start=1):
            aba.column_dimensions[get_column_letter(i)].width = larguras.get(chave, 15)
        for celula in aba[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1F4E78")
        for linha in aba.iter_rows(min_row=2):
            for celula in linha:
                celula.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(celula.value, str) and celula.value.startswith("http"):
                    celula.hyperlink = celula.value
                    celula.font = Font(color="0563C1", underline="single")
        aba.freeze_panes = "A2"
        aba.auto_filter.ref = aba.dimensions
    livro.save(caminho)


def _celula(valor):
    # Excel limita células a 32.767 caracteres e não aceita caracteres de controle.
    if isinstance(valor, str):
        valor = "".join(c for c in valor if c in "\t\n" or ord(c) >= 32)
        return valor[:32000]
    return valor


def _aba_resumo(aba, ocorrencias: list[LinhaOcorrencia], documentos: list[LinhaDocumento]) -> None:
    from collections import Counter

    from openpyxl.styles import Font

    aba.append(["Resumo da busca nas Reuniões Públicas da ANEEL"])
    aba["A1"].font = Font(bold=True, size=14)
    aba.append([])
    aba.append(["Documentos analisados", len(documentos)])
    aba.append(["Documentos com erro de leitura", sum(1 for d in documentos if d.situacao != "OK")])
    aba.append(["Total de ocorrências", len(ocorrencias)])
    aba.append([])
    aba.append(["Termo buscado", "Categoria", "Ocorrências", "Reuniões"])
    for celula in aba[aba.max_row]:
        celula.font = Font(bold=True)
    por_termo = Counter((o.termo, o.categoria) for o in ocorrencias)
    for (termo, categoria), total in por_termo.most_common():
        reunioes = sorted({o.reuniao for o in ocorrencias if o.termo == termo}, reverse=True)
        aba.append([termo, categoria, total, len(reunioes)])
    aba.column_dimensions["A"].width = 45
    aba.column_dimensions["B"].width = 16
    aba.column_dimensions["C"].width = 14
    aba.column_dimensions["D"].width = 12
