"""Extração de texto de HTML, PDF, DOCX e TXT."""

from __future__ import annotations

import io
import logging
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup

log = logging.getLogger(__name__)

EXTENSOES_DOCUMENTO = (".pdf", ".docx", ".doc", ".odt", ".rtf", ".txt", ".htm", ".html", ".xlsx", ".zip")


def tipo_por_conteudo(conteudo: bytes, content_type: str = "", nome: str = "") -> str:
    """Descobre o tipo do arquivo pelos primeiros bytes, cabeçalho HTTP ou extensão."""
    inicio = conteudo[:1024].lstrip()
    ct = (content_type or "").lower()
    nome = (nome or "").lower().split("?")[0]
    if inicio.startswith(b"%PDF") or "pdf" in ct:
        return "pdf"
    if inicio.startswith(b"PK"):
        if nome.endswith(".xlsx") or "spreadsheet" in ct:
            return "xlsx"
        if nome.endswith(".zip") or ("zip" in ct and "officedocument" not in ct):
            return "zip"
        return "docx"
    if inicio.startswith(b"\xd0\xcf\x11\xe0"):
        return "doc"
    if inicio.startswith(b"{\\rtf"):
        return "rtf"
    if "html" in ct or re.match(rb"(?is)\s*(<!--.*?-->\s*)*<(!doctype|html|head|body|table|div)", inicio):
        return "html"
    if nome.endswith(".txt") or "text/plain" in ct:
        return "txt"
    por_extensao = {".pdf": "pdf", ".docx": "docx", ".doc": "doc", ".rtf": "rtf",
                    ".xlsx": "xlsx", ".zip": "zip", ".htm": "html", ".html": "html"}
    for ext, tipo in por_extensao.items():
        if nome.endswith(ext):
            return tipo
    return "desconhecido"


def sopa(conteudo: bytes | str) -> BeautifulSoup:
    """Cria o BeautifulSoup. Com bytes, respeita o charset declarado (o site usa windows-1252)."""
    return BeautifulSoup(conteudo, "html.parser")


def texto_de_html(conteudo: bytes | str) -> str:
    doc = sopa(conteudo)
    for tag in doc(["script", "style", "noscript", "iframe", "form", "select", "head"]):
        tag.decompose()
    texto = doc.get_text("\n")
    return limpar_espacos(texto)


def texto_de_pdf(conteudo: bytes) -> str:
    from pypdf import PdfReader

    leitor = PdfReader(io.BytesIO(conteudo))
    if leitor.is_encrypted:
        try:
            leitor.decrypt("")
        except Exception:  # noqa: BLE001
            raise ValueError("PDF protegido por senha")
    paginas = []
    for i, pagina in enumerate(leitor.pages, start=1):
        try:
            paginas.append(pagina.extract_text() or "")
        except Exception as erro:  # noqa: BLE001
            log.warning("Falha ao ler a página %d do PDF: %s", i, erro)
    texto = limpar_espacos("\n".join(paginas))
    if not texto.strip():
        raise ValueError("PDF sem texto extraível (provavelmente digitalizado; seria preciso OCR)")
    return texto


def texto_de_docx(conteudo: bytes) -> str:
    import docx

    documento = docx.Document(io.BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    for tabela in documento.tables:
        for linha in tabela.rows:
            partes.append(" | ".join(c.text for c in linha.cells))
    return limpar_espacos("\n".join(partes))


def texto_de_xlsx(conteudo: bytes) -> str:
    import openpyxl

    livro = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    partes = []
    for planilha in livro.worksheets:
        for linha in planilha.iter_rows(values_only=True):
            valores = [str(v) for v in linha if v is not None]
            if valores:
                partes.append(" | ".join(valores))
    return limpar_espacos("\n".join(partes))


def texto_de_rtf(conteudo: bytes) -> str:
    texto = conteudo.decode("latin-1", errors="ignore")
    texto = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1252", "ignore"), texto)
    texto = re.sub(r"\\par[d]?", "\n", texto)
    texto = re.sub(r"\\[a-zA-Z]+-?\d* ?|[{}]", "", texto)
    return limpar_espacos(texto)


def decodificar_texto(conteudo: bytes) -> str:
    for codificacao in ("utf-8", "cp1252", "latin-1"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("latin-1", errors="ignore")


def extrair_texto(conteudo: bytes, content_type: str = "", nome: str = "") -> tuple[str, str]:
    """Retorna (tipo, texto). Lança ValueError se o formato não for suportado."""
    tipo = tipo_por_conteudo(conteudo, content_type, nome)
    if tipo == "pdf":
        return tipo, texto_de_pdf(conteudo)
    if tipo == "docx":
        return tipo, texto_de_docx(conteudo)
    if tipo == "xlsx":
        return tipo, texto_de_xlsx(conteudo)
    if tipo == "html":
        return tipo, texto_de_html(conteudo)
    if tipo == "rtf":
        return tipo, texto_de_rtf(conteudo)
    if tipo == "txt":
        return tipo, limpar_espacos(decodificar_texto(conteudo))
    if tipo == "zip":
        return tipo, texto_de_zip(conteudo)
    if tipo == "doc":
        raise ValueError("Formato .doc antigo não suportado; salve como .docx ou PDF")
    raise ValueError(f"Formato não reconhecido ({content_type or nome or 'sem tipo'})")


def texto_de_zip(conteudo: bytes) -> str:
    partes = []
    with zipfile.ZipFile(io.BytesIO(conteudo)) as arquivo_zip:
        for info in arquivo_zip.infolist():
            if info.is_dir() or not info.filename.lower().endswith(EXTENSOES_DOCUMENTO):
                continue
            try:
                _, texto = extrair_texto(arquivo_zip.read(info), nome=info.filename)
                partes.append(f"===== {info.filename} =====\n{texto}")
            except Exception as erro:  # noqa: BLE001
                log.warning("Não foi possível ler %s dentro do ZIP: %s", info.filename, erro)
    return "\n\n".join(partes)


def extrair_texto_arquivo(caminho: Path) -> tuple[str, str]:
    return extrair_texto(caminho.read_bytes(), nome=caminho.name)


def limpar_espacos(texto: str) -> str:
    texto = texto.replace("\xa0", " ").replace("\r", "\n")
    texto = re.sub(r"[ \t\f\v]+", " ", texto)
    texto = re.sub(r" *\n *", "\n", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()
